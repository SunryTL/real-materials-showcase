from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from real_workbench.app import create_app
from real_workbench.config import Settings
from real_workbench.packages import SEVEN_TABLE_NAMES


def make_client(tmp_path: Path) -> TestClient:
    settings = Settings(
        runtime_root=tmp_path / "runtime",
        pdf_inbox_root=tmp_path / "inbox",
        candidate_root=tmp_path / "candidates",
        paper_vault_root=tmp_path / "vault",
        research_root=tmp_path / "research",
        initial_admin_username="owner",
        initial_admin_password="correct horse battery staple",
        session_secret="test-secret-with-at-least-32-characters",
    )
    settings.pdf_inbox_root.mkdir(parents=True)
    return TestClient(create_app(settings))


def login(client: TestClient) -> None:
    response = client.post(
        "/api/v1/session/login",
        json={"username": "owner", "password": "correct horse battery staple"},
    )
    assert response.status_code == 200


def register_document(client: TestClient) -> int:
    response = client.post(
        "/api/v1/documents",
        files={"files": ("paper.pdf", b"%PDF-1.4\n10.1111/jace.20497\n%%EOF", "application/pdf")},
    )
    document_id = response.json()["items"][0]["document_id"]
    updated = client.patch(
        f"/api/v1/documents/{document_id}",
        json={
            "doi_confirmed": "10.1111/jace.20497",
            "target_family": "Lu-Gd solid solution",
            "target_fields": ["composition", "PL"],
        },
    )
    assert updated.status_code == 200
    return document_id


def test_codex_handoff_is_auditable_and_does_not_start_api_extraction(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    login(client)
    document_id = register_document(client)

    response = client.post(f"/api/v1/documents/{document_id}/codex-handoff")
    assert response.status_code == 200
    payload = response.json()
    assert payload["mode"] == "codex_handoff"
    assert payload["doi"] == "10.1111/jace.20497"
    assert len(payload["pdf_sha256"]) == 64
    assert "七表" in payload["task"]
    assert payload["authority_boundary"] == "candidate_only"


def test_candidate_workbook_import_requires_all_seven_tables_and_stays_candidate_only(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    login(client)
    document_id = register_document(client)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SEVEN_TABLE_NAMES:
        sheet = workbook.create_sheet(name)
        sheet.append(["sample_id"])
        if name == "sample_master":
            sheet.append(["candidate-sample-1"])
    path = tmp_path / "candidate.xlsx"
    workbook.save(path)

    response = client.post(
        "/api/v1/candidate-imports",
        data={"document_id": str(document_id)},
        files={"file": ("candidate.xlsx", path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "candidate_imported"
    assert payload["authority_write_enabled"] is False
    assert payload["sample_rows"] == 1
    assert isinstance(payload["job_id"], int)
    review = client.get(f"/api/v1/packages/{payload['job_id']}")
    assert review.status_code == 200
    assert review.json()["tables"]["sample_master"][0]["sample_id"] == "candidate-sample-1"

    incomplete = Workbook()
    incomplete.active.title = "sample_master"
    bad = tmp_path / "incomplete.xlsx"
    incomplete.save(bad)
    rejected = client.post(
        "/api/v1/candidate-imports",
        files={"file": ("incomplete.xlsx", bad.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert rejected.status_code == 422

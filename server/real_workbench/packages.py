from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


SEVEN_TABLE_NAMES = (
    "sample_master",
    "site_occupancy",
    "optical_measurement",
    "ceramic_process",
    "crystal_structure",
    "physical_descriptors",
    "data_source",
)

FIELD_EVIDENCE_COLUMNS = [
    "evidence_row_id", "table_name", "record_id", "field_name", "source_id",
    "provenance_type", "page_number", "figure_table_id", "unit", "uncertainty",
    "rule_version", "notes",
]
GROUPING_COLUMNS = [
    "sample_id", "proposed_group_series", "proposed_group_composition",
    "proposed_group_family", "reason", "evidence_source_id", "review_status",
]


def safe_doi(doi: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", doi).strip("_")


def _write_csv(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in columns})


def export_xlsx(package: Path, schemas: dict[str, list[str]]) -> None:
    workbook = Workbook()
    info = workbook.active
    info.title = "提取说明"
    info.append(["REAL文献候选提取包", "候选数据须经人工审核后才能进入正式数据库"])
    for table_name, columns in schemas.items():
        sheet = workbook.create_sheet(table_name)
        sheet.append(columns)
        source = package / ("tables" if table_name in SEVEN_TABLE_NAMES else "audit") / f"{table_name}.csv"
        with source.open("r", encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                sheet.append([row.get(column, "") for column in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B5CAD")
        sheet.freeze_panes = "A2"
    workbook.save(package / "extraction_review.xlsx")


def write_candidate_package(
    *,
    output_root: Path,
    schema_path: Path,
    doi: str,
    pdf_sha256: str,
    target_family: str,
    target_fields: list[str],
    payload: dict[str, Any],
    model: str,
    prompt_version: str,
) -> Path:
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    table_schemas: dict[str, list[str]] = schema["tables"]
    missing = set(SEVEN_TABLE_NAMES) - set(table_schemas)
    if missing:
        raise ValueError(f"七表schema不完整：{sorted(missing)}")
    package = output_root / safe_doi(doi)
    package.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone.utc).isoformat()
    metadata = {
        "schema_version": schema.get("schema_version", "1.0"),
        "doi": doi,
        "target_family": target_family,
        "target_fields": target_fields,
        "status": "candidate_ready",
        "pdf_sha256": pdf_sha256,
        "ai_model": model,
        "prompt_version": prompt_version,
        "generated_at": now,
        "authority_boundary": "candidate_only_never_overwrite_authority",
    }
    (package / "package.yaml").write_text(
        yaml.safe_dump(metadata, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    tables = payload.get("tables", {})
    for name in SEVEN_TABLE_NAMES:
        _write_csv(package / "tables" / f"{name}.csv", table_schemas[name], tables.get(name, []))
    _write_csv(package / "audit/field_evidence.csv", FIELD_EVIDENCE_COLUMNS, payload.get("field_evidence", []))
    _write_csv(package / "audit/grouping_review.csv", GROUPING_COLUMNS, payload.get("grouping_review", []))
    (package / "admission_review.md").write_text(
        "# 人工准入审核\n\n"
        f"AI候选摘要：{payload.get('admission_summary', '等待人工审核')}\n\n"
        "- [ ] 核心候选\n- [ ] 辅助数据\n- [ ] 排除\n- [ ] 待补证据\n",
        encoding="utf-8",
    )
    (package / "evidence_manifest.csv").write_text(
        "evidence_type,sha256,doi,notes\n"
        f"local_pdf,{pdf_sha256},{doi},PDF仅保存在本地且未复制到候选包\n",
        encoding="utf-8",
    )
    (package / "validation_report.json").write_text(
        json.dumps({"valid": False, "status": "pending_human_review", "errors": [], "warnings": ["尚未运行论文库完整校验器"]}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    export_xlsx(package, {**table_schemas, "field_evidence": FIELD_EVIDENCE_COLUMNS, "grouping_review": GROUPING_COLUMNS})
    return package


def read_package(package: Path) -> dict[str, Any]:
    metadata = yaml.safe_load((package / "package.yaml").read_text(encoding="utf-8"))
    tables: dict[str, list[dict[str, str]]] = {}
    for name in SEVEN_TABLE_NAMES:
        with (package / "tables" / f"{name}.csv").open("r", encoding="utf-8-sig", newline="") as stream:
            tables[name] = list(csv.DictReader(stream))
    return {"metadata": metadata, "tables": tables}


def import_candidate_bundle(
    *, output_root: Path, schema_path: Path, filename: str, data: bytes
) -> dict[str, Any]:
    """Stage a seven-table workbook/zip without touching the authority database."""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".zip"}:
        raise ValueError("只接受七表XLSX或包含七张CSV的ZIP")
    schema = json.loads(schema_path.read_text(encoding="utf-8"))["tables"]
    present: set[str] = set()
    headers: dict[str, list[str]] = {}
    row_counts: dict[str, int] = {}
    parsed_rows: dict[str, list[dict[str, Any]]] = {}
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        present = set(workbook.sheetnames)
        for name in SEVEN_TABLE_NAMES:
            if name not in present:
                continue
            rows = workbook[name].iter_rows(values_only=True)
            first = next(rows, ())
            headers[name] = [str(value).strip() for value in first if value is not None]
            parsed_rows[name] = [dict(zip(headers[name], row)) for row in rows if any(value not in (None, "") for value in row)]
            row_counts[name] = len(parsed_rows[name])
        workbook.close()
    else:
        try:
            archive = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile as error:
            raise ValueError("ZIP文件损坏或格式不正确") from error
        names = {Path(name).name: name for name in archive.namelist() if not name.endswith("/")}
        for table in SEVEN_TABLE_NAMES:
            member = names.get(f"{table}.csv")
            if not member:
                continue
            present.add(table)
            text = archive.read(member).decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            parsed_rows[table] = list(reader)
            headers[table] = [value.strip() for value in (reader.fieldnames or [])]
            row_counts[table] = len(parsed_rows[table])
        archive.close()
    missing_tables = sorted(set(SEVEN_TABLE_NAMES) - present)
    if missing_tables:
        raise ValueError(f"候选包缺少七表：{', '.join(missing_tables)}")
    missing_columns = {
        name: [column for column in schema[name] if column not in headers.get(name, [])]
        for name in SEVEN_TABLE_NAMES
    }
    digest = hashlib.sha256(data).hexdigest()
    target = output_root / "imports" / digest[:16]
    target.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name)
    (target / safe_name).write_bytes(data)
    for table in SEVEN_TABLE_NAMES:
        _write_csv(target / "tables" / f"{table}.csv", schema[table], parsed_rows.get(table, []))
    metadata = {
        "schema_version": "1.0", "status": "candidate_ready", "source": "codex_import",
        "import_sha256": digest, "authority_boundary": "candidate_only_never_overwrite_authority",
    }
    (target / "package.yaml").write_text(yaml.safe_dump(metadata, allow_unicode=True, sort_keys=False), encoding="utf-8")
    (target / "admission_review.md").write_text("# 人工准入审核\n\n- [ ] 核心候选\n- [ ] 辅助数据\n- [ ] 排除\n- [ ] 待补证据\n", encoding="utf-8")
    (target / "validation_report.json").write_text(json.dumps({"valid": False, "status": "pending_human_review", "missing_columns": missing_columns}, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "status": "candidate_imported",
        "authority_boundary": "candidate_only_never_overwrite_authority",
        "sha256": digest,
        "filename": safe_name,
        "tables": list(SEVEN_TABLE_NAMES),
        "row_counts": row_counts,
        "missing_columns": missing_columns,
        "package_path": target.relative_to(output_root).as_posix(),
    }
    (target / "import_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest
